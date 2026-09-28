import io
import re

import docx
import pdfplumber


class ExtractionError(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf(content)
    elif name.endswith(".docx"):
        text = _docx(content)
    elif name.endswith(".txt") or name.endswith(".md"):
        text = content.decode("utf-8", errors="ignore")
    else:
        raise ExtractionError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = text.strip()
    if len(text) < 50:
        raise ExtractionError(
            "Could not read text from this file. If it is a scanned image PDF, export it as a text PDF or DOCX."
        )
    return text


_BOLD_RE = re.compile(r"bold|black|heavy|semibold|demi", re.I)


def _is_bold(word: dict) -> bool:
    return bool(_BOLD_RE.search(word.get("fontname", "")))


def _pdf_lines(page) -> list[str]:
    """Rebuild visual lines: words grouped by vertical position, bold runs wrapped in **,
    and wide horizontal gaps (e.g. a right-aligned date) shown as ' | '."""
    words = page.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False)
    rows: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if rows and abs(rows[-1][0]["top"] - w["top"]) <= 3:
            rows[-1].append(w)
        else:
            rows.append([w])

    lines = []
    for row in rows:
        row.sort(key=lambda w: w["x0"])
        out, bold_open, prev = [], False, None
        for w in row:
            if prev is not None:
                gap = w["x0"] - prev["x1"]
                if gap > max(prev["size"], w["size"]) * 2.5:
                    if bold_open:
                        out.append("**")
                        bold_open = False
                    out.append(" | ")
                else:
                    out.append(" ")
            b = _is_bold(w)
            if b and not bold_open:
                out.append("**")
                bold_open = True
            elif not b and bold_open:
                # Close before the space we just added.
                if out and out[-1] == " ":
                    out.insert(len(out) - 1, "**")
                else:
                    out.append("**")
                bold_open = False
            out.append(w["text"])
            prev = w
        if bold_open:
            out.append("**")
        # PDF words split at font changes, so "**bold** ," needs its space removed.
        lines.append(re.sub(r"\*\* ([.,;:!?)])", r"**\1", "".join(out)))
    return lines


def _pdf(content: bytes) -> str:
    lines, links = [], []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            lines.extend(_pdf_lines(page))
            for link in page.hyperlinks or []:
                uri = link.get("uri")
                if uri and uri not in links:
                    links.append(uri)
    text = "\n".join(lines)
    if links:
        text += f"\n\n{_LINKS_HEADER} (for matching visible text to URLs only):\n" + "\n".join(links)
    return text


_WORD_RE = re.compile(r"[a-z0-9]+")
_LINKS_HEADER = "HYPERLINK TARGETS IN THE PDF"


def _words(text: str) -> list[str]:
    return _WORD_RE.findall(text.lower())


def _json_strings(value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [s for v in value.values() for s in _json_strings(v)]
    if isinstance(value, list):
        return [s for v in value for s in _json_strings(v)]
    return []


def missing_lines(source_text: str, data: dict) -> list[str]:
    """Source lines that mostly don't appear anywhere in the structured result.
    Compares consecutive word pairs, since single words repeat across a resume."""
    strings = _json_strings(data)
    have_words = set(_words(" ".join(strings)))
    have_pairs = {pair for s in strings for pair in zip(_words(s), _words(s)[1:])}
    missing = []
    for line in source_text.split(_LINKS_HEADER)[0].splitlines():
        words = _words(line)
        if not words:
            continue
        if len(words) == 1:
            ok = words[0] in have_words
        else:
            pairs = list(zip(words, words[1:]))
            ok = sum(p in have_pairs for p in pairs) / len(pairs) >= 0.6
        if not ok:
            missing.append(line.strip())
    return missing


def _docx(content: bytes) -> str:
    d = docx.Document(io.BytesIO(content))
    lines = []
    for p in d.paragraphs:
        parts = []
        for run in p.runs:
            t = run.text
            if run.bold and t.strip():
                t = f"**{t.strip()}** " if t.endswith(" ") else f"**{t.strip()}**"
            parts.append(t)
        lines.append("".join(parts))
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)
